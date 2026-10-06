from pathlib import Path

from hh_parser.export.xlsx import XLSXExporter
from hh_parser.db.repositories import EmployerRepository, VacancyRepository


class TestXLSXExporter:
    def test_export_creates_file(self, db_conn, sample_employer, tmp_path):
        emp_repo = EmployerRepository(db_conn)
        vac_repo = VacancyRepository(db_conn)
        emp_repo.upsert(sample_employer)

        output = tmp_path / "test_export.xlsx"
        exporter = XLSXExporter(emp_repo, vac_repo)
        path = exporter.export(output)

        assert path.exists()
        assert path.stat().st_size > 0

    def test_export_contains_employer(self, db_conn, sample_employer, tmp_path):
        emp_repo = EmployerRepository(db_conn)
        vac_repo = VacancyRepository(db_conn)
        emp_repo.upsert(sample_employer)

        output = tmp_path / "test_export.xlsx"
        exporter = XLSXExporter(emp_repo, vac_repo)
        exporter.export(output)

        from openpyxl import load_workbook
        wb = load_workbook(output)
        ws = wb["Работодатели"]
        assert ws.max_row >= 2
        assert ws.cell(row=2, column=2).value == "Тестовая организация"

    def test_export_empty_db(self, db_conn, tmp_path):
        emp_repo = EmployerRepository(db_conn)
        vac_repo = VacancyRepository(db_conn)

        output = tmp_path / "empty.xlsx"
        exporter = XLSXExporter(emp_repo, vac_repo)
        path = exporter.export(output)
        assert path.exists()
