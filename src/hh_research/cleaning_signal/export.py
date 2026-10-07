"""Export cleaning-signal results to XLSX with 6 sheets."""

from __future__ import annotations

import logging
from datetime import datetime, timezone
from pathlib import Path

from .aggregation import EmployerAggregate

logger = logging.getLogger(__name__)


def _employer_rows(employers: list[EmployerAggregate], enriched: dict[int, dict] | None = None):
    enriched = enriched or {}
    rows = []
    for emp in sorted(employers, key=lambda e: e.signal_score, reverse=True):
        e = enriched.get(emp.employer_id, {})
        rows.append({
            "Employer ID": emp.employer_id,
            "Employer": emp.employer_name,
            "Employer Category": emp.employer_category.value,
            "Region": " | ".join(emp.regions),
            "Total HH Vacancies": emp.total_hh_vacancies,
            "Relevant Vacancies": emp.total_relevant,
            "Cleaning Vacancies": emp.cleaning_count,
            "Facility Vacancies": emp.facility_count,
            "Janitorial Vacancies": emp.janitorial_count,
            "Unique Roles": emp.unique_role_count,
            "Unique Locations": emp.unique_location_count,
            "Fresh 7d": emp.fresh_7d_count,
            "Fresh 30d": emp.fresh_30d_count,
            "Latest Vacancy": emp.latest_relevant_publication or "",
            "Days Since Latest": emp.days_since_latest if emp.days_since_latest is not None else "",
            "Signal Score": emp.signal_score,
            "Signal Level": emp.signal_level,
            "Signal Reason": emp.signal_reason,
            "Cleaning Provider?": "Yes" if emp.is_cleaning_provider else "No",
            "HH URL": emp.employer_hh_url or "",
            "Official Site": e.get("official_website", emp.employer_site_url or ""),
            "Public Phone": e.get("official_phone", ""),
            "Public Email": e.get("official_email", ""),
            "Contact Confidence": e.get("contact_confidence", ""),
        })
    return rows


def _vacancy_rows(employers: list[EmployerAggregate]):
    rows = []
    for emp in employers:
        for v in emp.relevant_vacancies:
            rows.append({
                "Employer ID": emp.employer_id,
                "Employer": emp.employer_name,
                "Vacancy ID": v.vacancy_id,
                "Vacancy Name": v.name,
                "Vacancy URL": v.url or "",
                "Role Category": v.role_category.value,
                "Published At": v.published_at or "",
                "Area": v.area_name or "",
                "Salary From": v.salary_from if v.salary_from is not None else "",
                "Salary To": v.salary_to if v.salary_to is not None else "",
                "Currency": v.currency or "",
                "Matched Terms": ", ".join(v.matched_terms),
                "Relevance Score": round(v.relevance_score, 3),
            })
    return rows


def _high_signal_rows(employers: list[EmployerAggregate]):
    high = [e for e in employers if e.signal_level == "HIGH"]
    high.sort(key=lambda e: e.signal_score, reverse=True)
    rows = []
    for emp in high:
        rows.append({
            "Employer ID": emp.employer_id,
            "Employer": emp.employer_name,
            "Region": " | ".join(emp.regions),
            "Signal Score": emp.signal_score,
            "Relevant Vacancies": emp.total_relevant,
            "Fresh 7d": emp.fresh_7d_count,
            "Fresh 30d": emp.fresh_30d_count,
            "Unique Roles": emp.unique_role_count,
            "Unique Locations": emp.unique_location_count,
            "Cleaning Provider?": "Yes" if emp.is_cleaning_provider else "No",
            "Signal Reason": emp.signal_reason,
            "HH URL": emp.employer_hh_url or "",
        })
    return rows


def _enriched_contacts_rows(employers: list[EmployerAggregate], enriched: dict[int, dict]):
    rows = []
    for emp in employers:
        e = enriched.get(emp.employer_id, {})
        if not e:
            continue
        rows.append({
            "Employer ID": emp.employer_id,
            "Employer": emp.employer_name,
            "Signal Level": emp.signal_level,
            "Official Website": e.get("official_website", ""),
            "Public Phone": e.get("official_phone", ""),
            "Public Email": e.get("official_email", ""),
            "Official Address": e.get("official_address", ""),
            "Facility Phone": e.get("facility_contact_phone", ""),
            "Facility Email": e.get("facility_contact_email", ""),
            "Contact Confidence": e.get("contact_confidence", ""),
            "Enrichment Status": e.get("enrichment_status", ""),
        })
    return rows


def _cleaning_providers_rows(employers: list[EmployerAggregate]):
    providers = [e for e in employers if e.is_cleaning_provider]
    providers.sort(key=lambda e: e.signal_score, reverse=True)
    rows = []
    for emp in providers:
        rows.append({
            "Employer ID": emp.employer_id,
            "Employer": emp.employer_name,
            "Category": emp.employer_category.value,
            "Region": " | ".join(emp.regions),
            "Relevant Vacancies": emp.total_relevant,
            "Signal Score": emp.signal_score,
            "Signal Level": emp.signal_level,
            "Category Reason": emp.employer_category_reason,
            "HH URL": emp.employer_hh_url or "",
        })
    return rows


def _run_summary_rows(
    employers: list[EmployerAggregate],
    total_raw: int,
    total_relevant: int,
    total_rejected: int,
    enriched: dict[int, dict],
) -> list[dict]:
    end_customers = [e for e in employers if not e.is_cleaning_provider]
    cleaning_providers = [e for e in employers if e.is_cleaning_provider]
    high = [e for e in employers if e.signal_level == "HIGH"]
    medium = [e for e in employers if e.signal_level == "MEDIUM"]
    low = [e for e in employers if e.signal_level == "LOW"]

    enriched_count = len(enriched)
    with_phone = sum(1 for e in enriched.values() if e.get("official_phone"))
    with_email = sum(1 for e in enriched.values() if e.get("official_email"))
    with_facility = sum(
        1 for e in enriched.values()
        if e.get("facility_contact_phone") or e.get("facility_contact_email")
    )

    by_region: dict[str, int] = {}
    for emp in employers:
        for r in emp.regions:
            by_region[r] = by_region.get(r, 0) + 1

    by_category: dict[str, int] = {}
    for emp in employers:
        by_category[emp.employer_category.value] = by_category.get(emp.employer_category.value, 0) + 1

    by_role: dict[str, int] = {}
    for emp in employers:
        for v in emp.relevant_vacancies:
            by_role[v.role_category.value] = by_role.get(v.role_category.value, 0) + 1

    rows = [
        {"Metric": "Total raw vacancies reviewed", "Value": total_raw},
        {"Metric": "Relevant vacancies", "Value": total_relevant},
        {"Metric": "Rejected vacancies", "Value": total_rejected},
        {"Metric": "Unique employers", "Value": len(employers)},
        {"Metric": "End-customer candidates", "Value": len(end_customers)},
        {"Metric": "Cleaning providers", "Value": len(cleaning_providers)},
        {"Metric": "Unknown category", "Value": sum(1 for e in employers if e.employer_category.value == "OTHER")},
        {"Metric": "HIGH signal employers", "Value": len(high)},
        {"Metric": "MEDIUM signal employers", "Value": len(medium)},
        {"Metric": "LOW signal employers", "Value": len(low)},
        {"Metric": "Enriched employers", "Value": enriched_count},
        {"Metric": "Public phone coverage", "Value": with_phone},
        {"Metric": "Public email coverage", "Value": with_email},
        {"Metric": "Facility-specific contact coverage", "Value": with_facility},
        {"Metric": "", "Value": ""},
        {"Metric": "--- By Region ---", "Value": ""},
    ]
    for region, count in sorted(by_region.items(), key=lambda x: -x[1])[:20]:
        rows.append({"Metric": f"  {region}", "Value": count})

    rows.append({"Metric": "", "Value": ""})
    rows.append({"Metric": "--- By Employer Category ---", "Value": ""})
    for cat, count in sorted(by_category.items(), key=lambda x: -x[1]):
        rows.append({"Metric": f"  {cat}", "Value": count})

    rows.append({"Metric": "", "Value": ""})
    rows.append({"Metric": "--- By Role Category ---", "Value": ""})
    for role, count in sorted(by_role.items(), key=lambda x: -x[1]):
        rows.append({"Metric": f"  {role}", "Value": count})

    return rows


def export_cleaning_xlsx(
    employers: list[EmployerAggregate],
    total_raw: int,
    total_relevant: int,
    total_rejected: int,
    export_dir: str,
    enriched: dict[int, dict] | None = None,
    prefix: str = "cleaning-hiring-signal",
) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    out = Path(export_dir)
    out.mkdir(parents=True, exist_ok=True)
    ts = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    path = out / f"{prefix}.xlsx"
    enriched = enriched or {}

    wb = Workbook()

    ws_emp = wb.active
    ws_emp.title = "EMPLOYERS"
    emp_rows = _employer_rows(employers, enriched)
    if emp_rows:
        headers = list(emp_rows[0].keys())
        ws_emp.append(headers)
        for row in emp_rows:
            ws_emp.append([row[h] for h in headers])
        _style_sheet(ws_emp, headers)

    ws_vac = wb.create_sheet("RELEVANT VACANCIES")
    vac_rows = _vacancy_rows(employers)
    if vac_rows:
        headers = list(vac_rows[0].keys())
        ws_vac.append(headers)
        for row in vac_rows:
            ws_vac.append([row[h] for h in headers])
        _style_sheet(ws_vac, headers)

    ws_high = wb.create_sheet("HIGH SIGNAL")
    high_rows = _high_signal_rows(employers)
    if high_rows:
        headers = list(high_rows[0].keys())
        ws_high.append(headers)
        for row in high_rows:
            ws_high.append([row[h] for h in headers])
        _style_sheet(ws_high, headers)

    ws_enr = wb.create_sheet("ENRICHED CONTACTS")
    enr_rows = _enriched_contacts_rows(employers, enriched)
    if enr_rows:
        headers = list(enr_rows[0].keys())
        ws_enr.append(headers)
        for row in enr_rows:
            ws_enr.append([row[h] for h in headers])
        _style_sheet(ws_enr, headers)

    ws_cp = wb.create_sheet("CLEANING PROVIDERS")
    cp_rows = _cleaning_providers_rows(employers)
    if cp_rows:
        headers = list(cp_rows[0].keys())
        ws_cp.append(headers)
        for row in cp_rows:
            ws_cp.append([row[h] for h in headers])
        _style_sheet(ws_cp, headers)

    ws_sum = wb.create_sheet("RUN SUMMARY")
    sum_rows = _run_summary_rows(employers, total_raw, total_relevant, total_rejected, enriched)
    if sum_rows:
        headers = list(sum_rows[0].keys())
        ws_sum.append(headers)
        for row in sum_rows:
            ws_sum.append([row[h] for h in headers])
        header_font = Font(bold=True)
        for row in ws_sum.iter_rows(min_row=1, max_row=ws_sum.max_row, min_col=1, max_col=1):
            for cell in row:
                cell.font = header_font
        ws_sum.column_dimensions["A"].width = 45
        ws_sum.column_dimensions["B"].width = 60

    wb.save(str(path))
    logger.info("Cleaning signal XLSX exported: %s", path)
    return str(path)


def _style_sheet(ws, headers: list[str]):
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
    for col_idx in range(1, len(headers) + 1):
        max_len = max(
            len(str(ws.cell(row=r, column=col_idx).value or ""))
            for r in range(1, ws.max_row + 1)
        )
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 50)
    ws.auto_filter.ref = ws.dimensions
