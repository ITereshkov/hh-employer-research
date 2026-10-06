"""Export to CSV, XLSX (multi-sheet), JSON."""

from __future__ import annotations

import csv
import json
import logging
from datetime import datetime, timezone
from pathlib import Path

from .aggregation import EmployerRecord, SignalMetrics

logger = logging.getLogger(__name__)


def _employer_rows(records: list[EmployerRecord], metrics: dict[int, SignalMetrics],
                   enriched: dict[int, dict] | None = None):
    enriched = enriched or {}
    rows = []
    for r in records:
        m = metrics.get(r.employer_id, SignalMetrics())
        e = enriched.get(r.employer_id, {})
        rows.append({
            "employer_id": r.employer_id,
            "employer_name": r.employer_name,
            "employer_type": r.employer_type,
            "employer_type_confidence": r.employer_type_confidence,
            "classification_reason": r.classification_reason,
            "employer_hh_url": r.employer_hh_url or "",
            "area_id": r.area_id or "",
            "area_name": r.area_name or "",
            "industry": r.industry or "",
            "total_open_vacancies": r.total_open_vacancies,
            "relevant_open_vacancies": r.relevant_open_vacancies,
            "latest_relevant_vacancy_date": r.latest_relevant_vacancy_date or "",
            "oldest_current_relevant_vacancy_date": r.oldest_current_relevant_vacancy_date or "",
            "days_since_latest": r.days_since_latest if r.days_since_latest is not None else "",
            "relevant_vacancy_titles": " | ".join(r.relevant_vacancy_titles),
            "relevant_vacancy_urls": " | ".join(r.relevant_vacancy_urls),
            "active_relevant_count": m.active_relevant_count,
            "unique_relevant_roles": m.unique_relevant_roles,
            "fresh_signal": m.fresh_signal,
            "repeat_signal": m.repeat_signal,
            "multi_role_signal": m.multi_role_signal,
            "research_signal_score": m.research_signal_score,
            "score_explanation": m.score_explanation,
            "official_website": e.get("official_website", ""),
            "official_phone": e.get("official_phone", ""),
            "official_email": e.get("official_email", ""),
            "official_address": e.get("official_address", ""),
            "facility_contact_phone": e.get("facility_contact_phone", ""),
            "facility_contact_email": e.get("facility_contact_email", ""),
            "enrichment_signal_score": e.get("signal_score", ""),
        })
    return rows


def _vacancy_rows(records: list[EmployerRecord]):
    rows = []
    for r in records:
        for v in r.relevant_vacancies:
            rows.append({
                "employer_id": r.employer_id,
                "employer_name": r.employer_name,
                "employer_type": r.employer_type,
                "vacancy_id": v.vacancy_id,
                "vacancy_name": v.vacancy_name,
                "vacancy_url": v.vacancy_url or "",
                "published_at": v.published_at or "",
                "salary_from": v.salary_from if v.salary_from is not None else "",
                "salary_to": v.salary_to if v.salary_to is not None else "",
                "currency": v.currency or "",
                "schedule": v.schedule or "",
                "employment": v.employment or "",
                "experience": v.experience or "",
                "matched_terms": ", ".join(v.matched_terms),
                "match_reason": v.match_reason,
                "relevance_score": round(v.relevance_score, 3),
            })
    return rows


def _contact_rows(all_contacts: list[dict]) -> list[dict]:
    rows = []
    for c in all_contacts:
        rows.append({
            "employer_id": c.get("employer_id", ""),
            "employer_name": c.get("employer_name", ""),
            "contact_type": c.get("contact_type", ""),
            "value": c.get("value", ""),
            "source_url": c.get("source_url", ""),
            "source_type": c.get("source_type", ""),
            "confidence": c.get("confidence", 0.0),
            "is_primary": bool(c.get("is_primary", 0)),
            "is_facility": bool(c.get("is_facility", 0)),
            "retrieved_at": c.get("retrieved_at", ""),
        })
    return rows


def export_csv(records: list[EmployerRecord], metrics: dict[int, SignalMetrics],
               export_dir: str, prefix: str = "research",
               enriched: dict[int, dict] | None = None) -> tuple[str, str]:
    out = Path(export_dir)
    out.mkdir(parents=True, exist_ok=True)
    ts = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")

    emp_path = out / f"{prefix}_employers_{ts}.csv"
    emp_rows = _employer_rows(records, metrics, enriched)
    if emp_rows:
        keys = list(emp_rows[0].keys())
        with open(emp_path, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.DictWriter(f, fieldnames=keys)
            w.writeheader()
            w.writerows(emp_rows)

    vac_path = out / f"{prefix}_vacancies_{ts}.csv"
    vac_rows = _vacancy_rows(records)
    if vac_rows:
        keys = list(vac_rows[0].keys())
        with open(vac_path, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.DictWriter(f, fieldnames=keys)
            w.writeheader()
            w.writerows(vac_rows)

    logger.info("CSV exported: %s, %s", emp_path, vac_path)
    return str(emp_path), str(vac_path)


def export_json(records: list[EmployerRecord], metrics: dict[int, SignalMetrics],
                export_dir: str, prefix: str = "research",
                enriched: dict[int, dict] | None = None) -> str:
    out = Path(export_dir)
    out.mkdir(parents=True, exist_ok=True)
    ts = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    path = out / f"{prefix}_{ts}.json"
    enriched = enriched or {}

    data = []
    for r in records:
        m = metrics.get(r.employer_id, SignalMetrics())
        e = enriched.get(r.employer_id, {})
        data.append({
            "employer_id": r.employer_id,
            "employer_name": r.employer_name,
            "employer_type": r.employer_type,
            "employer_type_confidence": r.employer_type_confidence,
            "classification_reason": r.classification_reason,
            "employer_hh_url": r.employer_hh_url,
            "area_id": r.area_id,
            "area_name": r.area_name,
            "industry": r.industry,
            "total_open_vacancies": r.total_open_vacancies,
            "relevant_open_vacancies": r.relevant_open_vacancies,
            "latest_relevant_vacancy_date": r.latest_relevant_vacancy_date,
            "days_since_latest": r.days_since_latest,
            "contacts": {
                "website": e.get("official_website"),
                "phone": e.get("official_phone"),
                "email": e.get("official_email"),
                "address": e.get("official_address"),
                "facility_phone": e.get("facility_contact_phone"),
                "facility_email": e.get("facility_contact_email"),
            },
            "enrichment_signal_score": e.get("signal_score"),
            "relevant_vacancies": [
                {
                    "vacancy_id": v.vacancy_id,
                    "vacancy_name": v.vacancy_name,
                    "vacancy_url": v.vacancy_url,
                    "published_at": v.published_at,
                    "salary_from": v.salary_from,
                    "salary_to": v.salary_to,
                    "currency": v.currency,
                    "matched_terms": v.matched_terms,
                    "match_reason": v.match_reason,
                    "relevance_score": v.relevance_score,
                }
                for v in r.relevant_vacancies
            ],
            "signal_metrics": {
                "active_relevant_count": m.active_relevant_count,
                "days_since_latest": m.days_since_latest,
                "unique_relevant_roles": m.unique_relevant_roles,
                "fresh_signal": m.fresh_signal,
                "repeat_signal": m.repeat_signal,
                "multi_role_signal": m.multi_role_signal,
                "research_signal_score": m.research_signal_score,
                "score_explanation": m.score_explanation,
            },
        })

    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    logger.info("JSON exported: %s", path)
    return str(path)


def export_xlsx(records: list[EmployerRecord], metrics: dict[int, SignalMetrics],
                export_dir: str, run_summary: dict | None = None,
                prefix: str = "research",
                enriched: dict[int, dict] | None = None,
                all_contacts: list[dict] | None = None) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    out = Path(export_dir)
    out.mkdir(parents=True, exist_ok=True)
    ts = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    path = out / f"{prefix}_{ts}.xlsx"
    enriched = enriched or {}

    wb = Workbook()

    ws_emp = wb.active
    ws_emp.title = "Employers"
    emp_rows = _employer_rows(records, metrics, enriched)
    if emp_rows:
        headers = list(emp_rows[0].keys())
        ws_emp.append(headers)
        for row in emp_rows:
            ws_emp.append([row[h] for h in headers])
        _style_sheet(ws_emp, headers)

    ws_vac = wb.create_sheet("Relevant Vacancies")
    vac_rows = _vacancy_rows(records)
    if vac_rows:
        headers = list(vac_rows[0].keys())
        ws_vac.append(headers)
        for row in vac_rows:
            ws_vac.append([row[h] for h in headers])
        _style_sheet(ws_vac, headers)

    if all_contacts:
        ws_contacts = wb.create_sheet("Contacts")
        contact_rows = _contact_rows(all_contacts)
        if contact_rows:
            headers = list(contact_rows[0].keys())
            ws_contacts.append(headers)
            for row in contact_rows:
                ws_contacts.append([row[h] for h in headers])
            _style_sheet(ws_contacts, headers)

    ws_sum = wb.create_sheet("Run Summary")
    summary = run_summary or {}
    for k, v in summary.items():
        ws_sum.append([str(k), str(v)])
    header_font = Font(bold=True)
    for row in ws_sum.iter_rows(min_row=1, max_row=ws_sum.max_row, min_col=1, max_col=1):
        for cell in row:
            cell.font = header_font
    ws_sum.column_dimensions["A"].width = 40
    ws_sum.column_dimensions["B"].width = 60

    wb.save(str(path))
    logger.info("XLSX exported: %s", path)
    return str(path)


def _style_sheet(ws, headers: list[str]):
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
    for col_idx in range(1, len(headers) + 1):
        max_len = max(
            len(str(ws.cell(row=r, column=col_idx).value or ""))
            for r in range(1, ws.max_row + 1)
        )
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 50)
    ws.auto_filter.ref = ws.dimensions
