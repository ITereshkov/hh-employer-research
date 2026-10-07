from pathlib import Path
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from ..db.repositories import EmployerRepository, VacancyRepository, EmployerFilters


class XLSXExporter:
    EMPLOYER_COLUMNS = [
        ("ID", "id"),
        ("Название", "name"),
        ("Тип (индустрия)", "industry_name"),
        ("Регион", "area_name"),
        ("Релевантных вакансий", "vacancy_count"),
        ("Signal Score", "signal_score"),
        ("Официальный сайт", "official_website"),
        ("Телефон (общий)", "official_phone"),
        ("Email (общий)", "official_email"),
        ("Адрес", "official_address"),
        ("Роль (АХО)", "facility_contact_role"),
        ("Имя контакта (АХО)", "facility_contact_name"),
        ("Тел. АХО", "facility_contact_phone"),
        ("Email АХО", "facility_contact_email"),
        ("Источник контакта", "contact_source_url"),
        ("Confidence", "contact_confidence"),
        ("Статус обогащения", "enrichment_status"),
        ("HH URL", "alternate_url"),
        ("Первое обнаружение", "first_seen_at"),
        ("Обновлено", "last_updated_at"),
    ]

    VACANCY_COLUMNS = [
        ("ID", "id"),
        ("Название", "name"),
        ("Работодатель ID", "employer_id"),
        ("Регион", "area_name"),
        ("Категория", "category"),
        ("Ключевое слово", "keyword_match"),
        ("Зарплата от", "salary_from"),
        ("Зарплата до", "salary_to"),
        ("Валюта", "salary_currency"),
        ("Опубликовано", "published_at"),
        ("URL", "alternate_url"),
    ]

    def __init__(self, employer_repo: EmployerRepository, vacancy_repo: VacancyRepository):
        self._employer_repo = employer_repo
        self._vacancy_repo = vacancy_repo

    def export(
        self,
        output_path: Path,
        filters: Optional[EmployerFilters] = None,
    ) -> Path:
        output_path.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()

        employers = self._employer_repo.get_all(filters)
        self._write_employers_sheet(wb.active, employers)
        wb.active.title = "Работодатели"

        if employers:
            ws_vac = wb.create_sheet("Вакансии")
            self._write_vacancies_sheet(ws_vac, [e["id"] for e in employers])

            ws_sum = wb.create_sheet("Сводка")
            self._write_summary_sheet(ws_sum, employers)

        wb.save(output_path)
        return output_path

    def _write_employers_sheet(self, ws, employers: list[dict]) -> None:
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        for col_idx, (header, _) in enumerate(self.EMPLOYER_COLUMNS, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
            cell.border = thin_border

        for row_idx, emp in enumerate(employers, 2):
            for col_idx, (_, key) in enumerate(self.EMPLOYER_COLUMNS, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=emp.get(key))
                cell.border = thin_border

        ws.freeze_panes = "A2"
        for col_idx, (header, _) in enumerate(self.EMPLOYER_COLUMNS, 1):
            ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = max(
                len(header) * 1.5, 12
            )

    def _write_vacancies_sheet(self, ws, employer_ids: list[int]) -> None:
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")

        for col_idx, (header, _) in enumerate(self.VACANCY_COLUMNS, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill

        row = 2
        for eid in employer_ids:
            vacancies = self._vacancy_repo.get_by_employer(eid)
            for vac in vacancies:
                emp_vac_link = vac.copy()
                emp_vac_link["employer_id"] = eid
                for col_idx, (_, key) in enumerate(self.VACANCY_COLUMNS, 1):
                    ws.cell(row=row, column=col_idx, value=emp_vac_link.get(key))
                row += 1

        ws.freeze_panes = "A2"

    def _write_summary_sheet(self, ws, employers: list[dict]) -> None:
        header_font = Font(bold=True)

        ws.cell(row=1, column=1, value="Метрика").font = header_font
        ws.cell(row=1, column=2, value="Значение").font = header_font

        total = len(employers)
        with_website = sum(1 for e in employers if e.get("official_website"))
        with_phone = sum(1 for e in employers if e.get("official_phone"))
        with_email = sum(1 for e in employers if e.get("official_email"))
        with_facility = sum(1 for e in employers if e.get("facility_contact_phone"))
        enriched = sum(1 for e in employers if e.get("enrichment_status") == "done")

        rows = [
            ("Всего работодателей", total),
            ("Обогащено", enriched),
            ("С официальным сайтом", with_website),
            ("С телефоном", with_phone),
            ("С email", with_email),
            ("С контактом АХО", with_facility),
            ("Покрытие сайтом", f"{with_website/total*100:.1f}%" if total else "0%"),
            ("Покрытие телефоном", f"{with_phone/total*100:.1f}%" if total else "0%"),
            ("Покрытие АХО", f"{with_facility/total*100:.1f}%" if total else "0%"),
        ]

        for i, (metric, value) in enumerate(rows, 2):
            ws.cell(row=i, column=1, value=metric)
            ws.cell(row=i, column=2, value=value)

        ws.column_dimensions["A"].width = 30
        ws.column_dimensions["B"].width = 15
